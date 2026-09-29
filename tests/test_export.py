import io
import uuid
from unittest.mock import patch
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.db.database import SessionLocal
from app.db.models import (
    Candidate,
    CandidateProcessing,
    CandidateProcessingStatus,
    Domain,
    Evaluation,
    Hackathon,
    Resume,
    Screening,
    ScreeningResult,
    ScreeningResultStatus,
    ScreeningStatus,
)
from data_pipeline.exporter import export_screening_to_excel

client = TestClient(app)


def test_export_screening_to_excel_unit():
    metadata = {
        "screening_id": 101,
        "hackathon_name": "HackAI 2026",
        "domain_name": "Backend Engineering",
        "status": "COMPLETED",
        "selection_limit": 10,
        "waitlist_limit": 5,
        "total_candidates": 20,
        "processed_candidates": 19,
        "failed_candidates": 1,
        "selected_candidates": 10,
        "waitlisted_candidates": 5,
        "rejected_candidates": 4,
        "started_at": "2026-09-28T10:00:00Z",
        "completed_at": "2026-09-28T10:05:00Z",
    }
    results = [
        {
            "rank": 1,
            "decision": "SELECTED",
            "final_score": 92,
            "technical_skills_score": 18,
            "competitive_achievement_score": 14,
            "relevant_experience_score": 14,
            "projects_score": 24,
            "demonstrated_potential_score": 13,
            "domain_relevance_score": 9,
            "name": "Alice Smith",
            "email": "alice@example.com",
            "college": "MIT",
            "mobile": "+1234567890",
            "domain_name": "Backend Engineering",
            "linkedin_url": "https://linkedin.com/in/alice",
            "github_url": "https://github.com/alice",
            "resume_url": "https://example.com/alice.pdf",
            "processing_status": "COMPLETED",
            "failure_stage": "",
            "failure_reason": "",
        },
        {
            "rank": None,
            "decision": "FAILED",
            "final_score": None,
            "technical_skills_score": None,
            "competitive_achievement_score": None,
            "relevant_experience_score": None,
            "projects_score": None,
            "demonstrated_potential_score": None,
            "domain_relevance_score": None,
            "name": "Bob Stone",
            "email": "bob@example.com",
            "college": "Stanford",
            "mobile": "+1987654321",
            "domain_name": "Backend Engineering",
            "linkedin_url": "https://linkedin.com/in/bob",
            "github_url": "https://github.com/bob",
            "resume_url": "https://example.com/bob.pdf",
            "processing_status": "FAILED",
            "failure_stage": "EXTRACTION",
            "failure_reason": "Resume download failed: connection timeout",
        },
    ]

    excel_bytes = export_screening_to_excel(metadata, results)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    # Parse with openpyxl to verify workbook structure
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Screening Results" in wb.sheetnames
    assert "Summary" in wb.sheetnames

    ws_results = wb["Screening Results"]
    # Check headers
    headers = [cell.value for cell in ws_results[1]]
    assert "Rank" in headers
    assert "Decision" in headers
    assert "Final Score" in headers
    assert "Candidate Name" in headers
    assert "Email" in headers
    assert "Failure Reason" in headers

    # Check row 2 (Alice)
    row2 = [cell.value for cell in ws_results[2]]
    assert row2[0] == 1
    assert row2[1] == "SELECTED"
    assert row2[2] == 92
    assert "Alice Smith" in row2

    # Check row 3 (Bob)
    row3 = [cell.value for cell in ws_results[3]]
    assert row3[1] == "FAILED"
    assert "Bob Stone" in row3


def test_export_screening_api_endpoint():
    db = SessionLocal()
    hackathon = Hackathon(
        name=f"Export Test Hackathon {uuid.uuid4().hex}",
        description="Testing export endpoint",
    )
    db.add(hackathon)
    db.flush()

    domain = Domain(
        name="Cloud Systems",
        description="AWS, Docker, Kubernetes",
        hackathon_id=hackathon.id,
    )
    db.add(domain)
    db.flush()

    screening = Screening(
        hackathon_id=hackathon.id,
        domain_id=domain.id,
        selection_limit=5,
        waitlist_limit=3,
        status=ScreeningStatus.COMPLETED,
        total_candidates=1,
        processed_candidates=1,
        failed_candidates=0,
    )
    db.add(screening)
    db.flush()

    candidate = Candidate(
        name="Charlie Test",
        email=f"charlie.{uuid.uuid4().hex}@example.com",
        college="Tech Univ",
    )
    db.add(candidate)
    db.flush()

    resume = Resume(
        candidate_id=candidate.id,
        resume_url="https://example.com/charlie.pdf",
        file_type="pdf",
        file_size=1024,
    )
    db.add(resume)
    db.flush()

    processing = CandidateProcessing(
        screening_id=screening.id,
        candidate_id=candidate.id,
        status=CandidateProcessingStatus.COMPLETED,
    )
    db.add(processing)
    db.flush()

    evaluation = Evaluation(
        candidate_processing_id=processing.id,
        model="groq",
        prompt_version="v1",
        rubric_version="v1",
        evaluation_data={"sample": "data"},
        technical_skills_score=15,
        competitive_achievement_score=10,
        relevant_experience_score=12,
        projects_score=20,
        demonstrated_potential_score=11,
        domain_relevance_score=8,
        final_score=76,
    )
    db.add(evaluation)
    db.flush()

    screening_result = ScreeningResult(
        screening_id=screening.id,
        candidate_processing_id=processing.id,
        evaluation_id=evaluation.id,
        final_score=76,
        result_status=ScreeningResultStatus.SELECTED,
        rank=1,
    )
    db.add(screening_result)
    db.commit()

    screening_id = screening.id
    processing_id = processing.id
    candidate_id = candidate.id
    domain_id = domain.id
    hackathon_id = hackathon.id
    db.close()

    try:
        response = client.get(f"/api/v1/screenings/{screening_id}/export")
        assert response.status_code == 200
        assert (
            response.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "Content-Disposition" in response.headers
        assert f"screening_{screening_id}" in response.headers["Content-Disposition"]

        # Parse downloaded content
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        assert "Screening Results" in wb.sheetnames
        assert "Summary" in wb.sheetnames
    finally:
        db = SessionLocal()
        db.query(ScreeningResult).filter(ScreeningResult.screening_id == screening_id).delete()
        db.query(Evaluation).filter(Evaluation.candidate_processing_id == processing_id).delete()
        db.query(CandidateProcessing).filter(CandidateProcessing.screening_id == screening_id).delete()
        db.query(Screening).filter(Screening.id == screening_id).delete()
        db.query(Resume).filter(Resume.candidate_id == candidate_id).delete()
        db.query(Candidate).filter(Candidate.id == candidate_id).delete()
        db.query(Domain).filter(Domain.id == domain_id).delete()
        db.query(Hackathon).filter(Hackathon.id == hackathon_id).delete()
        db.commit()
        db.close()


def test_export_invalid_screening_id_returns_404():
    response = client.get("/api/v1/screenings/999999999/export")
    assert response.status_code == 404
    assert response.json()["detail"] == "Screening not found"
